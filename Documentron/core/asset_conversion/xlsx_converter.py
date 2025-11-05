"""
XLSX Converter

Converts Microsoft Excel spreadsheets (.xlsx) to normalized content.
"""

import hashlib
import zipfile
from pathlib import Path
from typing import Dict, Any, List
# Import openpyxl lazily in functions to avoid hard import dependency at module import time.

from .error_handling import MalformedFileError, EncryptedFileError, FileTooLargeError, DependencyError, report_skipped_file
from .provenance import get_library_versions

def calculate_file_hash(file_path: Path) -> str:
    """Calculate SHA256 hash of a file."""
    hash_sha256 = hashlib.sha256()
    with open(file_path, "rb") as f:
        for chunk in iter(lambda: f.read(4096), b""):
            hash_sha256.update(chunk)
    return hash_sha256.hexdigest()

def extract_xlsx_content(file_path: Path, max_size: int) -> Dict[str, Any]:
    """
    Extract content from an XLSX file.

    Args:
        file_path: Path to the XLSX file
        max_size: Maximum file size in bytes

    Returns:
        Dictionary with extracted content

    Raises:
        FileTooLargeError: If file exceeds max_size
        EncryptedFileError: If file is password-protected
        MalformedFileError: If file cannot be parsed
    """
    if file_path.stat().st_size > max_size:
        raise FileTooLargeError(f"File size {file_path.stat().st_size} exceeds limit {max_size}")

    try:
        from openpyxl import load_workbook  # type: ignore
    except ImportError as e:
        raise DependencyError(
            f"XLSX support requires openpyxl library; install with: pip install openpyxl"
        ) from e

    try:
        wb = load_workbook(filename=file_path, read_only=True, data_only=True)
    except zipfile.BadZipFile as e:
        raise EncryptedFileError(f"File appears to be encrypted: {e}")
    except Exception as e:
        # Check if it's an openpyxl-specific parsing error
        try:
            from openpyxl.utils.exceptions import InvalidFileException
            if isinstance(e, InvalidFileException):
                raise MalformedFileError(f"Cannot parse XLSX file: {e}")
        except ImportError:
            pass  # openpyxl not available, continue with generic handling
        # Fallback to generic error handling
        if "password" in str(e).lower() or "encrypted" in str(e).lower():
            raise EncryptedFileError(f"File appears to be encrypted: {e}")
        else:
            raise MalformedFileError(f"Cannot parse XLSX file: {e}")

    try:
        content = {
            "sheets": [],
            "metadata": {}
        }

        # Extract sheets
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            sheet_data = {
                "name": sheet_name,
                "rows": []
            }

            for row in ws.iter_rows(values_only=True):
                # Convert None to empty string, keep other values
                cleaned_row = [cell if cell is not None else "" for cell in row]
                if any(cleaned_row):  # Only include non-empty rows
                    sheet_data["rows"].append(cleaned_row)

            content["sheets"].append(sheet_data)

        # Extract basic metadata
        content["metadata"] = {
            "sheet_count": len(wb.sheetnames),
            "total_rows": sum(len(sheet["rows"]) for sheet in content["sheets"])
        }

        return content
    finally:
        wb.close()

def convert_xlsx_file(file_path: Path, output_dir: Path, max_size: int) -> Path:
    """
    Convert a single XLSX file to JSON artifact.

    Args:
        file_path: Path to the XLSX file
        output_dir: Output directory for artifacts
        max_size: Maximum file size

    Returns:
        Path to the created JSON artifact
    """
    try:
        file_hash = calculate_file_hash(file_path)
        content = extract_xlsx_content(file_path, max_size)

        provenance = {
            "converter_version": "1.0.0",
            "library_versions": get_library_versions()
        }

        from .writer import write_asset_artifact
        return write_asset_artifact(
            output_dir=output_dir,
            original_path=file_path,
            file_hash=file_hash,
            file_type="xlsx",
            content=content,
            provenance=provenance
        )

    except (FileTooLargeError, EncryptedFileError, MalformedFileError) as e:
        report_skipped_file(file_path, type(e).__name__.replace('Error', '').lower(), str(e))
        raise
