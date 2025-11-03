"""
DOCX Converter

Converts Microsoft Word documents (.docx) to normalized content.
"""

import hashlib
from pathlib import Path
from typing import Dict, Any, List
from docx import Document

from .error_handling import MalformedFileError, EncryptedFileError, FileTooLargeError, report_skipped_file
from .provenance import get_library_versions

def calculate_file_hash(file_path: Path) -> str:
    """Calculate SHA256 hash of a file."""
    hash_sha256 = hashlib.sha256()
    with open(file_path, "rb") as f:
        for chunk in iter(lambda: f.read(4096), b""):
            hash_sha256.update(chunk)
    return hash_sha256.hexdigest()

def extract_docx_content(file_path: Path, max_size: int) -> Dict[str, Any]:
    """
    Extract content from a DOCX file.

    Args:
        file_path: Path to the DOCX file
        max_size: Maximum file size in bytes

    Returns:
        Dictionary with extracted content

    Raises:
        FileTooLargeError: If file exceeds max_size
        EncryptedFileError: If file is password-protected
        MalformedFileError: If file cannot be parsed
    """
    if file_path.stat().st_size > max_size:
        raise FileTooLargeError(f"File size {file_path.stat().st_size} exceeds limit {max_size}")

    try:
        doc = Document(file_path)
    except Exception as e:
        if "password" in str(e).lower() or "encrypted" in str(e).lower():
            raise EncryptedFileError(f"File appears to be encrypted: {e}")
        else:
            raise MalformedFileError(f"Cannot parse DOCX file: {e}")

    content = {
        "title": "",
        "text_content": [],
        "tables": [],
        "metadata": {}
    }

    # Extract title (first paragraph if it's a heading)
    if doc.paragraphs:
        first_para = doc.paragraphs[0]
        if first_para.style and first_para.style.name.startswith('Heading'):
            content["title"] = first_para.text

    # Extract text content
    for para in doc.paragraphs:
        if para.text.strip():
            content["text_content"].append({
                "text": para.text,
                "style": para.style.name if para.style else "Normal"
            })

    # Extract tables
    for table in doc.tables:
        table_data = []
        for row in table.rows:
            row_data = [cell.text for cell in row.cells]
            table_data.append(row_data)
        content["tables"].append(table_data)

    # Extract basic metadata
    content["metadata"] = {
        "word_count": sum(len(para.text.split()) for para in doc.paragraphs),
        "paragraph_count": len(doc.paragraphs),
        "table_count": len(doc.tables)
    }

    return content

def convert_docx_file(file_path: Path, output_dir: Path, max_size: int) -> Path:
    """
    Convert a single DOCX file to JSON artifact.

    Args:
        file_path: Path to the DOCX file
        output_dir: Output directory for artifacts
        max_size: Maximum file size

    Returns:
        Path to the created JSON artifact
    """
    try:
        file_hash = calculate_file_hash(file_path)
        content = extract_docx_content(file_path, max_size)

        provenance = {
            "converter_version": "1.0.0",
            "library_versions": get_library_versions()
        }

        from .writer import write_asset_artifact
        return write_asset_artifact(
            output_dir=output_dir,
            original_path=file_path,
            file_hash=file_hash,
            file_type="docx",
            content=content,
            provenance=provenance
        )

    except (FileTooLargeError, EncryptedFileError, MalformedFileError) as e:
        report_skipped_file(file_path, type(e).__name__.replace('Error', '').lower(), str(e))
        raise